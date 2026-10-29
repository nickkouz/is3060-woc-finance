import sys
from docx import Document
from docx.shared import Pt, Inches
S='/home/claude/is3060-woc-finance/lab7-web-app/screenshots/'
OUTCOME=open(sys.argv[1]).read().strip() if len(sys.argv)>1 else "[PLACEHOLDER: learning outcome paragraph to be added]"
d=Document()
st=d.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(11)
d.add_heading('Lab 7: Building a Web Page with ChatGPT, Business Process Improvement',0)
d.add_paragraph('Nick Kouzmanoff | IS 3060-004 | Fall 2026')

d.add_heading('1. Process Overview',1)
for k,v in [('Organization or business area','Scholars Opportunity Fund (SOF), a student-run event-driven investment fund at the University of Utah. Sample content in the webpage is illustrative and does not describe official fund policy.'),
 ('Process','Research request intake and triage: an analyst asks a pod lead for research on an event-driven situation such as a merger, spin-off, or tender offer, and the fund produces a research memo.'),
 ('Users','Analysts who request research, pod leads who assign and review it, portfolio managers who use the memo, and the faculty advisor.'),
 ('Current problem','Requests arrive by message with missing details, status is hard to see, and two analysts can research the same company without knowing it.'),
 ('Better process','One structured intake form, one shared queue with visible status, duplicate warnings, and metrics that show workload and delays.')]:
    p=d.add_paragraph();p.add_run(k+': ').bold=True;p.add_run(v)

d.add_heading('2. Draft Comparison',1)
rows=[('Version','Main Features','Process Changes','Expected Process Benefit'),
('Draft 1: As-Is Process','Responsive Bootstrap page with navigation bar, hero, About the Process (users, inputs, activities, outputs, customers), current information, footer.','None. It documents the current message and spreadsheet process.','Makes the process visible and sets a baseline for finding weaknesses.'),
('Draft 2: Improved Process','Request form with required fields, format checks and a success message. Searchable queue with a status filter.','Structure: standard intake replaces free-form messages. Resources: one shared queue replaces scattered messages and a spreadsheet.','Fewer follow-up questions and incomplete requests. Analysts check status without interrupting pod leads.'),
('Draft 3: Innovative Process','Dashboard with KPIs and stage chart. Five-stage workflow with progress bars and a details modal. Duplicate-ticker warning. Priority recommendation tool.','Structure: defined stages and ownership of each step. Resources: data used for decisions. Integration: warnings connect separate requests.','Less duplicate work, earlier action on overdue items, and more consistent prioritization.')]
t=d.add_table(rows=len(rows),cols=4);t.style='Light Grid Accent 1'
for i,r in enumerate(rows):
    for j,c in enumerate(r):
        t.cell(i,j).text=c
        if i==0:
            for run in t.cell(i,j).paragraphs[0].runs: run.bold=True

d.add_heading('3. Screenshots',1)
for title,files in [('Draft 1: As-Is Process (index_v1.html)',['draft1_home.png','draft1_process.png']),
 ('Draft 2: Improved Process (index_v2.html)',['draft2_form_validation.png','draft2_queue_search.png']),
 ('Draft 3: Innovative Process (index_v3.html)',['draft3_dashboard.png','draft3_duplicate_priority.png','draft3_workflow_modal.png'])]:
    d.add_heading(title,2)
    caps={'draft1_home.png':'Home section and navigation bar.','draft1_process.png':'About the Process section with users, inputs, activities, outputs, and customers.',
    'draft2_form_validation.png':'New request form showing validation messages on an empty submit.','draft2_queue_search.png':'New request added to the queue and found with the search box.',
    'draft3_dashboard.png':'Dashboard with KPIs and requests by workflow stage.','draft3_duplicate_priority.png':'Duplicate-request warning and priority recommendation.','draft3_workflow_modal.png':'Details modal advancing a request to the next workflow stage.'}
    for f in files:
        d.add_picture(S+f,width=Inches(6.2));d.add_paragraph(caps[f]).italic=True

d.add_heading('4. AI Conversation Evidence',1)
d.add_paragraph('Disclosure: the three drafts in this submission were built with an AI coding assistant (Claude) and not with ChatGPT, and each feature was tested in a browser. No ChatGPT transcript exists for these drafts. The prompts below are the instructions used to build each draft, adapted from the lab prompts.')
d.add_heading('Draft 1 prompt',2)
d.add_paragraph('You are a web development assistant. Create a simple, responsive, single-page website using HTML, CSS, Bootstrap, and only client-side JavaScript. Organization or business area: Scholars Opportunity Fund, a student-run event-driven fund. Business process: research request intake and triage. Process users: analysts, pod leads, portfolio managers, faculty advisor. Current process problem: requests arrive by message with missing details, status is hard to see, and analysts duplicate work. Show the current or basic version of this process. Include a navigation bar with Home, About, and Contact; a hero section; an About the Process section; a Current Information section; a footer with contact or social links; and a responsive Bootstrap layout. Add comments in the code. Provide the complete code in one file named index_v1.html.')
d.add_heading('Draft 2 prompt',2)
d.add_paragraph('Using my previous webpage as the starting point, improve the business process and webpage functionality. Implement these two process improvements: (1) a request intake form with required fields, format validation, and a success message; (2) a searchable, filterable request queue showing status. Keep the site responsive with Bootstrap, use client-side HTML, CSS, and JavaScript only, label the new features in the code, explain how each improves the process, and do not remove the core information from Draft 1. Provide the complete code in one file named index_v2.html.')
d.add_heading('Draft 3 prompt',2)
d.add_paragraph('Improve the previous page into a connected research workflow, keeping everything in one client-side HTML file named index_v3.html. Add: (1) a dashboard with open, completed, overdue, and average-days-to-complete metrics plus a chart by stage, so managers can rebalance work early; (2) a five-stage workflow (Submitted, Triaged, In Research, In Review, Complete) with progress bars and a details modal that advances each request, so ownership and status are clear; (3) a duplicate-ticker warning on the form, so analysts do not repeat research; (4) a priority recommendation tool that scores event type and time left, so prioritization is consistent. Keep all earlier features and make everything work without a server.')

d.add_heading('5. Reflection',1)
refl=open('/home/claude/is3060-woc-finance/lab7-web-app/reflection.txt').read().strip().split('\n\n')
for para in refl: d.add_paragraph(para)

d.add_heading('6. Connection to Your Customized Learning Outcomes',1)
d.add_paragraph(OUTCOME)
d.save('/home/claude/is3060-woc-finance/lab7-web-app/Lab7_Kouzmanoff_Report.docx')
